"""Extract the source draft for targeted, XML-preserving review."""
from pathlib import Path
from docx import Document

folder = Path(__file__).resolve().parent
doc = Document(folder / "archive_before_renumbering/Lab08.docx")
lines = [f"P{i}: {p.text}" for i, p in enumerate(doc.paragraphs)]
for i, table in enumerate(doc.tables):
    lines.append(f"TABLE {i}")
    lines.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
(folder / "lab09_source_text.txt").write_text("\n".join(lines), encoding="utf-8")
print(f"Extracted {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables.")