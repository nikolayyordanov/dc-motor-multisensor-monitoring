"""Focused numerical, rendering and generated-document regression checks."""
from __future__ import annotations

import argparse
from copy import deepcopy
import csv
import json
from pathlib import Path
import struct
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

import matplotlib.pyplot as plt
from matplotlib.figure import Figure
import numpy as np
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image

import build_lab08_figures as builder
import plot_lab_current as shared

sys.path.insert(0, str(builder.ROOT / "paper"))
from build_dib_article import read_rigol_bin_segment, single_sided_amplitude


def write_bin(path, values, fs):
    path.parent.mkdir(parents=True, exist_ok=True)
    values = np.asarray(values, dtype="<f4")
    waveform = bytearray(60)
    struct.pack_into("<I", waveform, 0, 60)
    struct.pack_into("<II", waveform, 8, 1, len(values))
    struct.pack_into("<d", waveform, 32, 1 / fs)
    payload = values.tobytes()
    path.write_bytes(b"RG01" + struct.pack("<II", 84 + len(payload), 1)
                     + waveform + struct.pack("<IHHI", 12, 1, 4, len(payload)) + payload)


class NumericalTests(unittest.TestCase):
    def test_real_readers_and_fft_agree(self):
        for condition, _, _ in shared.CONDITIONS:
            with self.subTest(condition=condition):
                path = builder.ROOT / "DCData_mendeley/data" / condition / "current/speed030_current.bin"
                selected, fs, first = shared.read_interval(path, 2, 2)
                _, reference, ref_fs = read_rigol_bin_segment(path, 4.1)
                self.assertEqual(fs, ref_fs)
                np.testing.assert_array_equal(selected, reference[first:first + len(selected)])
                selected -= selected.mean()
                frequency, amplitude = shared.amplitude_spectrum(selected, fs)
                ref_frequency, ref_amplitude = single_sided_amplitude(selected, fs)
                np.testing.assert_array_equal(frequency, ref_frequency)
                np.testing.assert_array_equal(amplitude, ref_amplitude)
                self.assertAlmostEqual(fs / len(selected), 0.5, places=6)

    def test_fft_even_odd_dc_nyquist_and_sine(self):
        for n in (1000, 1001):
            y = np.ones(n) * 3
            _, amplitude = shared.amplitude_spectrum(y, 1000)
            self.assertAlmostEqual(amplitude[0], 3)
            window = np.hanning(n)
            expected = abs(np.fft.rfft(y * window)) / window.sum()
            expected[1:-1 if n % 2 == 0 else None] *= 2
            np.testing.assert_array_equal(amplitude, expected)
        n = 1000
        _, amplitude = shared.amplitude_spectrum((-1.0) ** np.arange(n) * 2, 1000)
        self.assertAlmostEqual(amplitude[-1], 2)
        _, amplitude = shared.amplitude_spectrum(4 * np.sin(2 * np.pi * 100 * np.arange(n) / n), 1000)
        self.assertAlmostEqual(amplitude[100], 4, places=5)

    def test_envelope_covers_every_sample_and_final_extreme(self):
        y = np.sin(np.arange(100003))
        y[-1] = 50
        centers, low, high = shared.overview_envelope(y, 1000)
        edges = np.linspace(0, len(y), 2001, dtype=int)
        self.assertEqual(len(centers), 2000)
        for i, (a, b) in enumerate(zip(edges[:-1], edges[1:])):
            self.assertEqual(low[i], y[a:b].min())
            self.assertEqual(high[i], y[a:b].max())
        self.assertEqual(high[-1], 50)
        for n in (1, 7, 2000):
            self.assertEqual(len(shared.overview_envelope(np.arange(n), 1)[0]), n)

    def test_invalid_interval_and_truncated_file(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "test.bin"
            write_bin(path, np.arange(100), 100)
            for start, duration in ((-1, 1), (0, 0), (0, 2), (float("nan"), 1)):
                with self.assertRaises(ValueError):
                    shared.read_interval(path, start, duration)
            path.write_bytes(path.read_bytes()[:-1])
            with self.assertRaises(ValueError):
                shared.read_interval(path, 0, 0.5)

    def test_layout_full_interval_scale_original_metrics_and_return(self):
        with tempfile.TemporaryDirectory() as folder:
            dataset = Path(folder)
            values = np.sin(np.arange(20000) / 5)
            values[-1] = 100  # Outside the 60-ms detail; must set the common scale.
            for condition, _, _ in shared.CONDITIONS:
                write_bin(dataset / "data" / condition / "current/speed030_current.bin", values, 10000)
            figures = []
            def capture(fig, *_args, **_kwargs):
                if fig not in figures:
                    figures.append(fig)
            args = argparse.Namespace(dataset=dataset, speed=30, start=0, duration=2,
                                      window_ms=60, max_hz=1500, output=dataset / "out")
            with patch.object(Figure, "savefig", capture):
                provenance = shared.generate(args)
            time_fig, fft_fig = figures
            self.assertEqual(len(time_fig.axes), 8)
            self.assertEqual(len(fft_fig.axes), 4)
            expected = values.astype("<f4").astype(float)
            expected -= expected.mean()
            limit = abs(expected).max() * 1.08
            for index, ax in enumerate(time_fig.axes):
                np.testing.assert_allclose(ax.get_ylim(), [-limit, limit])
                self.assertEqual(ax.get_xlim(), (0, 2 if index % 2 == 0 else 60))
                if index % 2:
                    np.testing.assert_array_equal(ax.lines[0].get_ydata(), expected[:601])
                else:
                    self.assertEqual(len(ax.collections[0].get_segments()), 2000)
            with (args.output / "example_metrics.csv").open(encoding="utf-8-sig") as fh:
                metrics = list(csv.DictReader(fh))
            self.assertAlmostEqual(float(metrics[0]["rms_ac_V"]), np.sqrt(np.mean(expected ** 2)))
            self.assertEqual(provenance["records"][0]["samples"], 20000)
            self.assertEqual(provenance["commanded_speed_rpm"], 600)
            self.assertIn("display only", provenance["overview_rendering"])
            for fig in figures:
                plt.close(fig)


class DocumentTests(unittest.TestCase):
    def test_split_runs_and_math_remain_in_place(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("Неоптимална настройка; оптимал").bold = True
        p.add_run("на настройка ")
        math = OxmlElement("m:oMath")
        math_run = OxmlElement("m:r")
        math_text = OxmlElement("m:t")
        math_text.text = "x"
        math_run.append(math_text)
        math.append(math_run)
        p._p.append(math)
        p.add_run(" след формула")
        original = deepcopy(p._p)
        self.assertTrue(builder.replace_text(p, "оптимална настройка", "номинална настройка"))
        self.assertEqual(p.text, "Неоптимална настройка; номинална настройка  след формула")
        self.assertEqual(builder.canonical_xml(math), builder.canonical_xml(original.xpath(".//m:oMath")[0]))
        self.assertEqual(builder.text_only_skeleton(original), builder.text_only_skeleton(p._p))
        self.assertEqual(p.runs[-1].text, " след формула")

    def test_locked_revision_uses_new_version(self):
        doc = unittest.mock.Mock()
        doc.save.side_effect = [PermissionError(), None]
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "Lab08_revised.docx"
            target.with_stem("Lab08_revised_v2").touch()
            self.assertEqual(builder.save_revision(doc, target).stem, "Lab08_revised_v3")

    def test_generated_docx_and_artifacts(self):
        report = json.loads((builder.OUT / "document_validation.json").read_text(encoding="utf-8"))
        docpath = builder.FOLDER / report["output"]
        doc = Document(docpath)
        source = Document(builder.FOLDER / "Lab08.docx")
        self.assertEqual(len(doc.tables), len(source.tables) + 1)
        self.assertEqual(len(doc.inline_shapes), len(source.inline_shapes) + 2)
        self.assertGreater(report["formula_count_preserved"], 0)
        self.assertEqual([etree.tostring(m) for m in source.element.body.xpath(".//m:oMath")],
                         [etree.tostring(m) for m in doc.element.body.xpath(".//m:oMath")])
        text = "\n".join(p.text for p in doc.paragraphs)
        for prefix, caption in (("Фиг. 8.5.", builder.CAPTION_TIME), ("Фиг. 8.6.", builder.CAPTION_FFT)):
            self.assertEqual([p.text for p in doc.paragraphs if p.text.startswith(prefix)], [caption])
        for phrase in ("зададени 600", "не е потвърдено", "18 периода", "не е независимо повторно",
                       "честотната стъпка са:", "котвения ток", "постояннотоковите системи"):
            self.assertIn(phrase, text)
        self.assertIn(source.paragraphs[141].text, text)  # Original conditional AC question.
        gains = next(t for t in doc.tables if "P05.04" in t.cell(0, 1).text)
        self.assertEqual([[c.text for c in r.cells][1:] for r in gains.rows[1:]],
                         [["10", "2"], ["10", "2"], ["0,2", "2"], ["10", "0,1"]])
        for name, checksum in report["protected_sha256"].items():
            self.assertEqual(builder.digest(builder.FOLDER / name), checksum)
        setting = doc.settings.element.find(qn("w:doNotAutoCompressPictures"))
        self.assertEqual(setting.get(qn("w:val")), "true")
        with zipfile.ZipFile(docpath) as archive:
            media = [archive.read(n) for n in archive.namelist() if n.startswith("word/media/")]
            for stem in ("fig8_5_current_waveforms", "fig8_6_current_spectra"):
                self.assertIn((builder.OUT / f"{stem}.png").read_bytes(), media)
        for out in (builder.OUT, builder.FOLDER / "figures_lab08_shared"):
            provenance = json.loads((out / "figure_provenance.json").read_text(encoding="utf-8"))
            for r in provenance["records"]:
                self.assertEqual(builder.digest(Path(provenance["dataset_root"]) / r["source"]), r["sha256"])
                self.assertEqual(r["overview_bins"], 2000)
            for stem in ("fig8_5_current_waveforms", "fig8_6_current_spectra"):
                for ext in ("png", "pdf", "svg"):
                    self.assertGreater((out / f"{stem}.{ext}").stat().st_size, 1000)
                with Image.open(out / f"{stem}.png") as image:
                    self.assertAlmostEqual(image.info["dpi"][0], 600, places=1)
        self.assertEqual((builder.OUT / "example_metrics.csv").read_bytes(),
                         (builder.FOLDER / "figures_lab08_shared/example_metrics.csv").read_bytes())


if __name__ == "__main__":
    unittest.main(verbosity=2)