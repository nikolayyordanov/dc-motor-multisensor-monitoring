"""Generate shared real-data figures and a separate, XML-verified Lab08 revision."""
from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from lxml import etree

from plot_lab_current import generate as generate_shared

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT / "exercises"
OUT = FOLDER / "figures_lab08"


def digest(path):
    with path.open("rb") as fh:
        return hashlib.file_digest(fh, "sha256").hexdigest()


def generate():
    return generate_shared(argparse.Namespace(
        dataset=ROOT / "DCData_mendeley", speed=30, start=2.0, duration=2.0,
        window_ms=60.0, max_hz=1500.0, output=OUT))


def replace_text(paragraph, old, new):
    """Replace across w:t/run boundaries, never reconstruct math or paragraph XML.

    Replacement inherits the first affected run's formatting. Unaffected text,
    run properties, drawings, fields and OMML elements remain untouched.
    """
    nodes = paragraph._p.xpath(".//w:t")
    text = "".join(node.text or "" for node in nodes)
    match = re.search(r"(?<!\w)" + re.escape(old), text)
    if match is None:
        return False
    start = match.start()
    end = start + len(old)
    offset = 0
    inserted = False
    for node in nodes:
        value = node.text or ""
        right = offset + len(value)
        if offset < end and right > start:
            left_text = value[:max(0, start - offset)]
            right_text = value[max(0, end - offset):]
            node.text = left_text + (new if not inserted else "") + right_text
            node.set(qn("xml:space"), "preserve")
            inserted = True
        offset = right
    return True


def canonical_xml(element):
    """Ignore unused inherited namespace declarations, not elements/attributes."""
    return etree.tostring(element, method="c14n", exclusive=True)


def text_only_skeleton(element):
    """Structure/properties with only intentionally editable w:t content masked."""
    clone = deepcopy(element)
    for node in clone.xpath(".//w:t"):
        node.text = "TEXT"
        node.attrib.pop(qn("xml:space"), None)
    return canonical_xml(clone)


def add_after(anchor, text):
    element = OxmlElement("w:p")
    anchor.addnext(element)
    paragraph = Paragraph(element, None)
    paragraph.add_run(text)
    return element


def save_revision(doc, target):
    """A locked Office file is never force-closed or overwritten indirectly."""
    version = 1
    while True:
        candidate = target if version == 1 else target.with_stem(f"{target.stem}_v{version}")
        try:
            doc.save(candidate)
            return candidate
        except PermissionError:
            version += 1
            # Do not overwrite another already existing version when falling back.
            while target.with_stem(f"{target.stem}_v{version}").exists():
                version += 1


CAPTION_TIME = (
    "Фиг. 8.5. Примерни времедиаграми при задание за скорост 30%, без реверсиране: "
    "а) номинални настройки; б) разхлабен фундамент; в) неоптимален РС; "
    "г) неоптимален РТ при номинален РС. Всеки ред показва целия интервал 2–4 s "
    "вляво и първите 60 ms (2,000–2,060 s от записа) вдясно. Общият изглед е "
    "min/max обвивка в 2000 интервала само за визуализация; детайлът е с оригиналните "
    "проби. Средното е отстранено за целите 2 s; общият амплитуден мащаб се определя "
    "от всички проби на четирите интервала, без нормировка по максимум.")
CAPTION_FFT = (
    "Фиг. 8.6. Едностранни амплитудни спектри за целите интервали 2–4 s при задание "
    "30%, без реверсиране: а) номинални настройки; б) разхлабен фундамент; "
    "в) неоптимален РС; г) неоптимален РТ при номинален РС. Прозорец на Хан с "
    "корекция чрез сумата на прозореца; честотна стъпка ≈ 0,5 Hz; без усредняване "
    "и допълване с нули. Пунктираните линии са при 300, 600, 900 и 1200 Hz; "
    "честотният и амплитудният мащаб са еднакви за четирите записа.")


def insert(provenance):
    source = FOLDER / "Lab08.docx"
    protected = [source, FOLDER / "Lab08_with_figures.docx"]
    protected_hashes = {path: digest(path) for path in protected if path.exists()}
    dataset = Path(provenance["dataset_root"])
    assert all(digest(dataset / r["source"]) == r["sha256"] for r in provenance["records"])
    doc = Document(source)
    original_shapes = len(doc.inline_shapes)
    # Retain references to EVERY original body child, including tables/section XML.
    originals = [(node, deepcopy(node)) for node in doc.element.body]
    paragraphs = [Paragraph(node, doc._body) for node in doc.element.body.xpath(".//w:p")]
    paragraph_originals = {p._p: deepcopy(p._p) for p in paragraphs}
    paragraph_ids = {p._p: i for i, p in enumerate(paragraphs)}
    changed = set()
    image_paragraphs = set()

    def edit(p, old, new):
        if replace_text(p, old, new):
            changed.add(p._p)

    # Never replace the substring inside НЕоптимална: these remain distinct conditions.
    for p in paragraphs:
        for old, new in (("Оптимална настройка", "Номинална настройка"),
                         ("оптимална настройка", "номинална настройка"),
                         ("постояннотокови системите", "постояннотоковите системи"),
                         ("Едностранният амплитуден спектър и разделителната способност по честота са:",
                          "Едностранният амплитуден спектър и честотната стъпка са:"),
                         ("разделителната способност е около 0,05 Hz", "честотната стъпка е около 0,05 Hz")):
            edit(p, old, new)
        if p.text.startswith("1. Кои са основните източници"):
            edit(p, "котвеният ток", "котвения ток")
        if "При използваното AC свързване" in p.text:
              edit(p, "При използваното AC свързване на входа се потиска постоянната съставка "
                  "и се отслабват ниските честоти според честотната характеристика на входа.",
                 "Запишете действителното AC/DC свързване на канала. За предоставените записи "
                 "то не е потвърдено. Численото отстраняване на средната стойност не доказва "
                 "AC свързване и не възстановява евентуално отслабените ниски честоти.")
        if p.text.startswith("Фиг. 8.5."):
            edit(p, p.text, CAPTION_TIME)
        if p.text.startswith("Фиг. 8.6."):
            edit(p, p.text, CAPTION_FFT)

    def anchor(prefix):
        found = [p._p for p in doc.paragraphs if p.text.lstrip().startswith(prefix)]
        assert len(found) == 1, (prefix, len(found))
        return found[0]

    add_after(anchor("Спектралният анализ"),
              "Честотната стъпка е Δf = fs/N = 1/T: при T = 20 s тя е 0,05 Hz, "
              "а при T = 2 s е 0,5 Hz. Това е разстояние между FFT коефициентите, "
              "не пълната способност за разграничаване на близки съставки. Прозорецът на Хан "
              "намалява спектралното разтичане, но разширява главния лоб. Корекцията на "
              "кохерентното усилване се прави чрез сумата на прозореца; удвояват се само "
              "вътрешните положителни честотни коефициенти, без DC и Nyquist (ако е наличен). "
              "В примерите няма усредняване или допълване с нули.")
    add_after(anchor("Обработката се извършва"),
              "Примерите могат да се възпроизведат със самостоятелния скрипт plot_lab_current.py "
              "и указанията README_figures_bg.md, споделяни към Приложение 2; не е необходим "
              "достъп до останалите програми на проекта. FFT и всички показатели, включително "
              "RMS, се изчисляват от оригиналните проби на целия избран 2-секунден интервал.")
    add_after(anchor("Сравнението се провежда"),
              "Заданието 30% от номиналните 2000 min⁻¹ означава зададени 600 min⁻¹, "
              "а не измерена скорост. Само ако действителната скорост е 600 min⁻¹, един "
              "оборот трае 100 ms. Избраният 60 ms детайл съдържа 18 периода на 300 Hz "
              "и не обхваща пълен оборот при тази условна скорост.")
    gains_anchor = add_after(anchor("При изследване на разхлабения фундамент"),
              "Базови коефициенти за сравняваните условия (параметри P05.04 за РС и "
              "P06.02 за РТ); при разхлабен фундамент се запазват номиналните усилвания:")
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ("Състояние", "РС — P05.04", "РТ — P06.02")):
        cell.text = text
    for values in (("Номинални настройки", "10", "2"),
                   ("Разхлабен фундамент", "10", "2"),
                   ("Неоптимален РС", "0,2", "2"),
                   ("Неоптимален РТ при номинален РС", "10", "0,1")):
        for cell, text in zip(table.add_row().cells, values):
            cell.text = text
    gains_anchor.addnext(table._tbl)
    add_after(anchor("За всяко от условията"),
              "Фиг. 8.5 използва два времеви мащаба: общ изглед на 2–4 s и детайл на "
              "първите 60 ms. За общия изглед всички проби се обобщават с min/max обвивка "
              "в 2000 интервала, само за визуализация и компактен векторен файл. Детайлът "
              "не се обобщава. Еднаквите ординати се определят от пълните 2 s за всички "
              "условия. Обвивката не се използва за FFT или количествените показатели.")
    add_after(anchor("Получените резултати се сравняват"),
              "Задачи за анализа: изчислете процентната промяна на RMS спрямо номиналния "
              "режим: 100·(RMSусловие/RMSбаза − 1), ако RMSбаза ≠ 0. Сравнете съставките "
              "около 300, 600 и 900 Hz при еднаква методика. Повторете обработката за втори "
              "неприпокриващ се 2-секунден интервал (например 4–6 s), за да оцените "
              "стабилността в рамките на записа; това не е независимо повторно измерване. "
              "Обосновете изводите с данните, без да се изисква предварително зададено "
              "разграничаване на състоянията. Липса на разлика в тока не доказва изправност; "
              "самостоятелният токов анализ има ограничения.")

    replacements = {"[Място за времедиаграмите при четирите условия.]": "fig8_5_current_waveforms.png",
                    "[Място за амплитудните спектри при четирите условия.]": "fig8_6_current_spectra.png"}
    section = doc.sections[0]
    width = min(int(section.page_width - section.left_margin - section.right_margin), int(Inches(6.5)))
    for p in paragraphs:
        if p.text.strip() in replacements:
            image = OUT / replacements[p.text.strip()]
            p.clear()  # Only the two known plain placeholders; never formula paragraphs.
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(image), width=width)
            image_paragraphs.add(p._p)
    assert len(image_paragraphs) == 2
    setting = doc.settings.element.find(qn("w:doNotAutoCompressPictures"))
    if setting is None:
        setting = OxmlElement("w:doNotAutoCompressPictures")
        doc.settings.element.append(setting)
    setting.set(qn("w:val"), "true")

    # Compare original nodes by identity, not shifted paragraph/table positions.
    # Even edited paragraphs/tables must retain ALL non-text XML (including math).
    for p in paragraphs:
        old = paragraph_originals[p._p]
        if p._p in image_paragraphs:
            assert not old.xpath(".//m:oMath | .//w:drawing | .//w:object")
        elif p._p in changed:
            assert text_only_skeleton(old) == text_only_skeleton(p._p)
        else:
            assert canonical_xml(old) == canonical_xml(p._p), "Untouched original paragraph changed"
    for node, old in originals:
        if node in image_paragraphs:
            continue
        affected = node in changed or any(p in changed for p in node.iter())
        if affected:
            assert text_only_skeleton(old) == text_only_skeleton(node), "Non-text XML changed"
        else:
            assert canonical_xml(old) == canonical_xml(node), "Untouched original XML changed"
    math_before = [canonical_xml(m) for _, old in originals for m in old.xpath(".//m:oMath")]
    math_after = [canonical_xml(m) for m in doc.element.body.xpath(".//m:oMath")]
    assert math_before == math_after, "Formula XML changed"

    target = save_revision(doc, FOLDER / "Lab08_revised.docx")
    check = Document(target)
    assert len(check.inline_shapes) == original_shapes + 2
    assert len(check.tables) == len(Document(source).tables) + 1
    # Round trip must preserve the entire revised body, not merely extracted text.
    assert doc.element.body.xml == check.element.body.xml, "DOCX round-trip changed XML"
    assert all(digest(path) == checksum for path, checksum in protected_hashes.items())
    assert all(digest(dataset / r["source"]) == r["sha256"] for r in provenance["records"])
    report = {"output": target.name,
              "protected_sha256": {p.name: checksum for p, checksum in protected_hashes.items()},
              "changed_original_paragraph_ids": sorted(paragraph_ids[p] for p in changed),
              "image_original_paragraph_ids": sorted(paragraph_ids[p] for p in image_paragraphs),
              "original_body_children_verified": len(originals),
              "formula_count_preserved": len(math_before),
              "original_tables_preserved_except_tracked_text_edits": len(Document(source).tables),
              "added_tables": 1, "dataset_sha256_verified": True}
    (OUT / "document_validation.json").write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Saved {target.name}; preserved {len(math_before)} formula objects and all untouched original XML.")
    return target


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--document-only", action="store_true", help="Reuse existing figures and checked provenance")
    args = parser.parse_args()
    provenance = (json.loads((OUT / "figure_provenance.json").read_text(encoding="utf-8"))
                  if args.document_only else generate())
    insert(provenance)