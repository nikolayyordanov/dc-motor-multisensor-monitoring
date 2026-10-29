"""Convert a read-only legacy Word draft to a separate working DOCX."""
from pathlib import Path
import hashlib
import win32com.client
from docx import Document

folder = Path(__file__).resolve().parent
source = folder / "archive_before_renumbering/Lab09.doc"
before = hashlib.sha256(source.read_bytes()).hexdigest()
target = folder / "Lab08_working.docx"
word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
word.DisplayAlerts = 0
try:
    doc = word.Documents.Open(str(source), ReadOnly=True, AddToRecentFiles=False)
    try:
        doc.SaveAs2(str(target), FileFormat=16, AddToRecentFiles=False)
    finally:
        doc.Close(SaveChanges=0)
finally:
    word.Quit()
assert hashlib.sha256(source.read_bytes()).hexdigest() == before
doc = Document(target)
text = "\n".join(f"{i}: {p.text}" for i, p in enumerate(doc.paragraphs))
text += "\nTABLES\n" + "\n".join(
    f"Table {i}: " + "\n".join(" | ".join(c.text for c in r.cells) for r in t.rows)
    for i, t in enumerate(doc.tables))
(folder / "lab08_source_text.txt").write_text(text, encoding="utf-8")
print(f"Converted read-only source; paragraphs={len(doc.paragraphs)}, tables={len(doc.tables)}, images={len(doc.inline_shapes)}")