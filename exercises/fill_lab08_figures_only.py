"""Insert only the two yellow figure requests; preserve all other source bytes."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import hashlib
import json
import re
import zipfile

from docx import Document
from docx.shared import Inches
from lxml import etree

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "Lab08_working.docx"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
REQUESTS = {"10.3": "fig8_3_features.png", "10.4": "fig8_4_confusion.png"}


def canonical(node):
    return etree.tostring(node, method="c14n", exclusive=True)


def main():
    checksum = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    with zipfile.ZipFile(SOURCE) as archive:
        infos = archive.infolist()
        original = {item.filename: archive.read(item) for item in infos}
    updated = dict(original)
    document = original["word/document.xml"]
    root = etree.fromstring(document)
    rels = etree.fromstring(original["word/_rels/document.xml.rels"])
    rel_ids = {node.get("Id") for node in rels}
    drawing_id = max([int(n.get("id")) for n in root.xpath("//wp:docPr", namespaces=NS)] + [0]) + 1
    replacements = []
    for match in re.finditer(rb"<w:p\b[^>]*>.*?</w:p>", document, re.DOTALL):
        # Parse in the original namespace context without changing the source XML.
        fragment = match.group()
        wrapper = etree.Element("wrapper", nsmap=root.nsmap)
        wrapper_xml = etree.tostring(wrapper)
        parsed = etree.fromstring(wrapper_xml[:-2] + b">" + fragment + b"</wrapper>")[0]
        text = "".join(parsed.xpath(".//w:t/text()", namespaces=NS))
        if not text.startswith("[Място за фигура"):
            continue
        key = next((key for key in REQUESTS if f"фигура {key}" in text), None)
        if key is None:
            continue
        assert parsed.xpath('.//w:highlight[@w:val="yellow"]', namespaces=NS)
        assert not parsed.xpath(".//w:object | .//w:drawing", namespaces=NS)
        image_path = HERE / "figures_lab08" / REQUESTS[key]
        image_name = "yellow_request_" + REQUESTS[key]
        assert f"word/media/{image_name}" not in original
        rel_id = "rIdYellowFigure" + key.replace(".", "_")
        assert rel_id not in rel_ids
        temporary = Document()
        run = temporary.add_paragraph().add_run()
        run.add_picture(str(image_path), width=Inches(6.3))
        new_run = deepcopy(run._r)
        new_run.xpath(".//a:blip")[0].set("{" + NS["r"] + "}embed", rel_id)
        new_run.xpath(".//wp:docPr")[0].set("id", str(drawing_id))
        drawing_id += 1
        for child in list(parsed):
            if child.tag != "{" + NS["w"] + "}pPr":
                parsed.remove(child)
        parsed.append(new_run)
        replacements.append((fragment, etree.tostring(parsed), key))
        updated[f"word/media/{image_name}"] = image_path.read_bytes()
        etree.SubElement(rels, "{" + REL + "}Relationship", Id=rel_id,
                         Type=NS["r"] + "/image", Target="media/" + image_name)
    assert {key for _, _, key in replacements} == set(REQUESTS)
    for before, after, _ in replacements:
        assert document.count(before) == 1
        document = document.replace(before, after, 1)
    updated["word/document.xml"] = document
    # Keep relationship bytes unchanged except for the two appended entries.
    extra = b"".join(etree.tostring(node) for node in list(rels)[-2:])
    original_rels = original["word/_rels/document.xml.rels"]
    updated["word/_rels/document.xml.rels"] = original_rels.replace(b"</Relationships>", extra + b"</Relationships>")
    types = original["[Content_Types].xml"]
    if not etree.fromstring(types).xpath('//*[local-name()="Default" and @Extension="png"]'):
        types = types.replace(b"</Types>", b'<Default Extension="png" ContentType="image/png"/></Types>')
    updated["[Content_Types].xml"] = types
    target = HERE / "Lab08_figures_only.docx"
    version = 1
    while True:
        try:
            with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
                for item in infos:
                    archive.writestr(item, updated[item.filename])
                for name in updated.keys() - original.keys():
                    archive.writestr(name, updated[name])
            break
        except PermissionError:
            version += 1
            target = HERE / f"Lab08_figures_only_v{version}.docx"
    with zipfile.ZipFile(target) as archive:
        result = {name: archive.read(name) for name in archive.namelist()}
    restored = result["word/document.xml"]
    for before, after, _ in replacements:
        assert restored.count(after) == 1
        restored = restored.replace(after, before, 1)
    assert restored == original["word/document.xml"], "Unrelated document XML changed"
    allowed = {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
    assert all(result[name] == data for name, data in original.items() if name not in allowed)
    check_root = etree.fromstring(result["word/document.xml"])
    for query in ("//w:tbl", "//w:object"):
        assert [canonical(n) for n in root.xpath(query, namespaces=NS)] == [canonical(n) for n in check_root.xpath(query, namespaces=NS)]
    check_rels = etree.fromstring(result["word/_rels/document.xml.rels"])
    assert len(check_rels) == len(etree.fromstring(original_rels)) + 2
    assert len(Document(target).inline_shapes) == len(Document(SOURCE).inline_shapes) + 2
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == checksum
    report = {"source": SOURCE.name, "output": target.name, "source_sha256": checksum,
              "changed_placeholders": [key for _, _, key in replacements],
              "all_other_document_xml_byte_identical": True,
              "tables_and_legacy_objects_unchanged": True,
              "other_existing_package_parts_byte_identical": True,
              "captions_numbering_and_text_unchanged": True}
    (HERE / "figures_only_validation.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    print("Saved", target.name, "- only two yellow figure placeholders replaced; all other content verified unchanged.")


if __name__ == "__main__":
    main()