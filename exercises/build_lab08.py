"""Reproducible teaching examples and separate revised Lab08 DOCX.

Run after inspect_lab08.py converts the archived read-only legacy DOC.
No source measurement or original manuscript is modified.
"""
from pathlib import Path
import argparse
import hashlib
import json
import re
import shutil

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
import numpy as np
import pandas as pd
import soundfile as sf
from scipy.signal import periodogram
from scipy.stats import kurtosis
from sklearn.base import clone
from sklearn.ensemble import RandomForestClassifier
from sklearn.metrics import accuracy_score, f1_score, confusion_matrix, ConfusionMatrixDisplay
from sklearn.model_selection import StratifiedGroupKFold
from sklearn.neighbors import KNeighborsClassifier
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from sklearn.svm import SVC
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from build_lab09_figures import canonical_xml, save_revision, replace_text

ROOT = Path(__file__).resolve().parents[1]
HERE = Path(__file__).resolve().parent
OUT = HERE / "figures_lab08"
ORIGINALS = HERE / "archive_before_renumbering"
DATA = ROOT / "DCData_mendeley"
CONDITIONS = ["normal_no_reversal", "loose_foundation_no_reversal",
              "suboptimal_control_no_reversal", "suboptimal_control_rt_no_reversal"]
LABELS = ["Номинално", "Разхлабен\nфундамент", "Неоптимален\nРС", "Неоптимален\nРТ"]
BANDS = [(0, 100), (100, 250), (250, 500), (500, 1000),
         (1000, 2000), (2000, 5000), (5000, 10000)]


def sha(path):
    with path.open("rb") as fh:
        return hashlib.file_digest(fh, "sha256").hexdigest()


def export(fig, name):
    fig.tight_layout()
    for extension in ("png", "pdf", "svg"):
        fig.savefig(OUT / f"{name}.{extension}", dpi=600, bbox_inches="tight")
    plt.close(fig)


def diagram(name, labels):
    fig, ax = plt.subplots(figsize=(10, 7))
    ax.set(xlim=(0, 10), ylim=(0, len(labels) * 1.3))
    ax.axis("off")
    for i, text in enumerate(labels):
        y = (len(labels) - i - 1) * 1.3 + .2
        ax.add_patch(FancyBboxPatch((.4, y), 9.2, .85, boxstyle="round,pad=.12",
                                   facecolor="#edf4fb", edgecolor="#30658b"))
        ax.text(5, y + .425, text, ha="center", va="center", fontsize=10)
        if i:
            ax.annotate("", (5, y + 1), (5, y + 1.3), arrowprops={"arrowstyle": "->"})
    export(fig, name)


def experiment():
    meta = pd.read_csv(DATA / "metadata.csv")
    meta = meta[(meta.sensor == "sound_vibrometer") & meta.condition.isin(CONDITIONS)]
    speeds = sorted(set.intersection(*(set(meta[meta.condition == c].speed_percent.astype(int))
                                       for c in CONDITIONS)))
    meta = meta[meta.speed_percent.isin(speeds)].sort_values(["condition", "speed_percent"])
    rows, sources = [], []
    for _, row in meta.iterrows():
        path = DATA / row.new_path
        checksum = sha(path)
        samples, fs = sf.read(path, dtype="float64", always_2d=True)
        if fs != 44100 or samples.shape[1] < 1:
            raise ValueError(f"Unexpected WAV layout: {path}")
        signal = samples[fs:, 0]  # Omit first second; use first channel, not stereo mixing.
        n = round(.5 * fs)
        for i in range(len(signal) // n):
            y = signal[i*n:(i+1)*n].copy()
            y -= y.mean()
            rms = float(np.sqrt(np.mean(y*y)))
            if rms == 0 or not np.isfinite(y).all():
                raise ValueError(f"Degenerate window: {path}, {i}")
            freq, power = periodogram(y, fs, window="hann", detrend=False, scaling="density")
            total = power.sum()
            values = [rms, float(np.ptp(y)), float(np.max(np.abs(y))/rms),
                      float(kurtosis(y, fisher=False, bias=True))]
            values += [float(power[(freq >= lo) & (freq < hi)].sum()/total) for lo, hi in BANDS]
            rows.append({"record": row.new_path, "speed": int(row.speed_percent),
                         "class": CONDITIONS.index(row.condition), "window": i,
                         **dict(zip(["RMS", "PTP", "CF", "Kurtosis"] +
                                    [f"band_{lo}_{hi}" for lo, hi in BANDS], values))})
        assert sha(path) == checksum
        sources.append({"path": row.new_path, "sha256": checksum})
    frame = pd.DataFrame(rows)
    columns = ["RMS", "PTP", "CF", "Kurtosis"] + [f"band_{lo}_{hi}" for lo, hi in BANDS]
    X, y = frame[columns].to_numpy(), frame["class"].to_numpy()
    folds = list(StratifiedGroupKFold(5, shuffle=True, random_state=42).split(X, y, frame.record))
    models = {"kNN": make_pipeline(StandardScaler(), KNeighborsClassifier(5)),
              "SVM": make_pipeline(StandardScaler(), SVC(C=1, gamma="scale", kernel="rbf")),
              "Random Forest": RandomForestClassifier(n_estimators=200, min_samples_leaf=2,
                                                       class_weight="balanced", random_state=42, n_jobs=-1)}
    metrics, predictions = [], {}
    for group_name, indices in (("Времева област", range(4)), ("Честотна област", range(4, 11)),
                                ("Времева и честотна област", range(11))):
        for name, prototype in models.items():
            pred = np.full(len(y), -1)
            for fold, (train, test) in enumerate(folds):
                assert not set(frame.record.iloc[train]) & set(frame.record.iloc[test])
                model = clone(prototype).fit(X[train][:, indices], y[train])
                pred[test] = model.predict(X[test][:, indices])
                if group_name == "Времева и честотна област":
                    frame.loc[test, "fold"] = fold
            votes = pd.DataFrame({"record": frame.record, "true": y, "pred": pred}).groupby("record")
            truth = votes["true"].first().to_numpy()
            result = votes["pred"].agg(lambda v: int(np.bincount(v, minlength=4).argmax())).to_numpy()
            metrics.append({"features": group_name, "model": name,
                            "accuracy_record": accuracy_score(truth, result),
                            "macro_f1_record": f1_score(truth, result, average="macro", zero_division=0),
                            "accuracy_window": accuracy_score(y, pred),
                            "macro_f1_window": f1_score(y, pred, average="macro", zero_division=0)})
            if group_name == "Времева и честотна област":
                predictions[name] = (truth, result)
                frame[f"prediction_{name}"] = pred
    frame.to_csv(OUT / "features_predictions.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(metrics).to_csv(OUT / "classification_metrics.csv", index=False, encoding="utf-8-sig")
    # One marker per independent source recording; window distributions are not repetitions.
    means = frame.groupby(["record", "class"])[["RMS", "Kurtosis"]].mean().reset_index()
    fig, ax = plt.subplots(figsize=(8, 5))
    for c, label in enumerate(LABELS):
        subset = means[means["class"] == c]
        ax.scatter(subset.RMS, subset.Kurtosis, label=label.replace("\n", " "), s=40)
    ax.set(xlabel="Среден RMS на прозорците (цифров пълен обхват)", ylabel="Среден куртозис (Pearson)")
    ax.grid(alpha=.25)
    ax.legend()
    export(fig, "fig8_3_features")
    fig, axes = plt.subplots(1, 3, figsize=(13, 4))
    for ax, (name, (truth, pred)) in zip(axes, predictions.items()):
        ConfusionMatrixDisplay(confusion_matrix(truth, pred, labels=range(4)),
                               display_labels=LABELS).plot(ax=ax, colorbar=False, cmap="Blues")
        ax.set_title(name)
        ax.set_xlabel("Предсказан клас")
        ax.set_ylabel("Действителен клас")
        ax.tick_params(axis="x", labelsize=7)
        ax.tick_params(axis="y", labelsize=7)
    export(fig, "fig8_4_confusion")
    provenance = {"speeds": speeds, "recordings": len(meta), "windows": len(frame),
                  "channel": 1, "skip_s": 1, "window_s": .5, "overlap": 0,
                  "features": columns, "spectral_method": "periodic Hann periodogram; relative band power / total power through Nyquist; half-open bands",
                  "validation": "5-fold StratifiedGroupKFold, group=source recording, seed=42; fixed hyperparameters; no tuning",
                  "primary_metrics": "pooled out-of-fold recording majority vote; ties choose lowest class index",
                  "models": {k: str(v) for k, v in models.items()}, "sources": sources}
    (OUT / "provenance.json").write_text(json.dumps(provenance, indent=2, ensure_ascii=False), encoding="utf-8")
    return metrics, provenance


def reuse_results():
    """Renumber the existing experiment without retraining any classifier.

    Parallel floating-point prediction can resolve exact RF voting ties
    differently. A numbering-only change must retain the approved results.
    """
    source = ORIGINALS / "figures_lab09"
    info = json.loads((source / "provenance.json").read_text(encoding="utf-8"))
    assert all(sha(DATA / row["path"]) == row["sha256"] for row in info["sources"])
    for name in ("classification_metrics.csv", "features_predictions.csv", "provenance.json"):
        shutil.copy2(source / name, OUT / name)
    for stem in ("3_features", "4_confusion"):
        for extension in ("png", "pdf", "svg"):
            shutil.copy2(source / f"fig9_{stem}.{extension}", OUT / f"fig8_{stem}.{extension}")
    return pd.read_csv(OUT / "classification_metrics.csv").to_dict("records"), info


def revise(metrics, info):
    original = ORIGINALS / "Lab09.doc"
    checksum = sha(original)
    doc = Document(HERE / "Lab08_working.docx")
    math_before = [canonical_xml(m) for m in doc.element.body.xpath(".//m:oMath")]
    objects_before = [canonical_xml(m) for m in doc.element.body.xpath(".//w:object")]
    # Number labels in text nodes only; preserve equation and drawing objects.
    for node in doc.element.body.xpath(".//w:t | .//m:t"):
        if node.text:
            node.text = re.sub(r"\b(?:9|10)\.(\d+)", r"8.\1", node.text)
            node.text = node.text.replace("УПРАЖНЕНИЕ № 9", "УПРАЖНЕНИЕ № 8")
            node.text = node.text.replace("lab10.py", "build_lab08.py")
            node.text = node.text.replace("аписите са групирани", "Записите са групирани")
            node.text = node.text.replace("упражнения № 9 и № 10", "упражнение № 9")
    # Legacy Word often splits heading/figure numbers between runs.
    for element in doc.element.body.xpath(".//w:p"):
        paragraph = Paragraph(element, doc._body)
        for match in list(re.finditer(r"\b(?:9|10)\.(\d+)", paragraph.text)):
            replace_text(paragraph, match.group(), "8." + match.group(1))
    images = ["fig8_1_pipeline", "fig8_2_dataset", "fig8_3_features", "fig8_4_confusion"]
    captions = [
        "Фиг. 8.1. Обработка и оценяване без пренасяне на информация: разделяне по записи; стандартизиране само в обучаващата част.",
        "Фиг. 8.2. Организация на набора: отделни, несинхронни измервания при съответстващи работни условия.",
        "Фиг. 8.3. RMS и куртозис (Pearson): една точка е средната стойност на прозорците от един запис. Диаграмата е описателна, не доказателство за качество на класификацията.",
        "Фиг. 8.4. Матрици на объркванията за kNN, SVM и Random Forest: обединени извънобучаващи предсказвания, по един глас на запис; еднакви пет групови дяла за трите модела."]
    for p in list(doc.paragraphs):
        if "[Място за фигура" in p.text:
            p.clear()
        for i in range(4):
            if p.text.startswith(f"Фиг. 8.{i+1}."):
                p.text = captions[i]
                element = OxmlElement("w:p")
                p._p.addprevious(element)
                image_p = Paragraph(element, p._parent)
                image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                image_p.paragraph_format.keep_with_next = True
                image_p.add_run().add_picture(str(OUT / f"{images[i]}.png"), width=Inches(6.3))
                for run in p.runs:
                    run.font.size = Pt(9)
    notes = [
        f"Примерен възпроизводим учебен експеримент: {info['recordings']} WAV записа и {info['windows']} прозореца; общи задания {', '.join(map(str, info['speeds']))}%. Използва се първият канал, без смесване на стереоканалите; пропуска се 1 s и се извличат неприпокриващи се прозорци по 0,5 s. От всеки прозорец се изважда средното. Амплитудата е спрямо цифровия пълен обхват, не в калибрирани вибрационни единици.",
        "Спектралните признаци са относителна мощност в седемте ленти, получена от периодограма с периодичен прозорец на Хан. Използват се полуотворени граници [долна, горна), без двойно броене. Знаменателят е мощността в целия диапазон до Nyquist, включително над 10 kHz; затова сборът на седемте признака може да е под 1. Куртозисът е Pearson (нормално разпределение: 3), без корекция за крайна извадка.",
        "Моделите използват предварително фиксирани настройки: kNN k=5; SVM RBF, C=1, gamma=scale; Random Forest 200 дървета, минимум 2 образеца в лист, балансирани класови тегла, seed=42. StandardScaler и kNN/SVM са в общ Pipeline. Няма подбор на параметри върху оценъчните данни; при настройване е необходима вътрешна групова проверка.",
        "Таблиците и матриците показват резултати на ниво запис: класът се избира чрез мнозинство на прозоречните предсказвания (при равенство — най-малкият индекс). Всички прозорци от един запис са в един дял. Сегментите не са независими повторни измервания. Това е оценка върху задържани записи от една установка, не доказателство за пренос към нов двигател или сесия.",
        "При многосензорно разширение групирайте по работна точка (състояние и задание), като всички модалности остават заедно. Отделно изследвайте задържане на цяло задание за скорост; текущото разделяне по записи не гарантира обобщаване към невиждана скорост. Клипването на WAV и вариациите между прозорци могат да влияят на признаците.",
        "Задачи: сравнете временни, спектрални и комбинирани признаци; обяснете кои класове се бъркат; сравнете прозоречните и записните показатели от CSV; проверете влиянието на дължината на прозореца. Не избирайте настройки според най-добрия резултат върху същите оценъчни дялове. Значимостта от случайна гора не е причинно обяснение; при допълнителен анализ използвайте пермутационна значимост върху задържани данни."]
    anchor = next(p for p in doc.paragraphs if "Контролни въпроси" in p.text)
    for text in notes:
        element = OxmlElement("w:p")
        anchor._p.addprevious(element)
        Paragraph(element, anchor._parent).add_run(text)
    for row in doc.tables[1].rows[1:]:
        metric = next(m for m in metrics if m['model'] == row.cells[0].text and m['features'] == "Времева и честотна област")
        row.cells[1].text = f"{100*metric['accuracy_record']:.1f}"
        row.cells[2].text = f"{metric['macro_f1_record']:.3f}"
    for row in doc.tables[2].rows[1:]:
        for cell, name in zip(row.cells[1:], ("kNN", "SVM", "Random Forest")):
            metric = next(m for m in metrics if m['model'] == name and m['features'] == row.cells[0].text)
            cell.text = f"{metric['macro_f1_record']:.3f}"
    setting = doc.settings.element.find(qn("w:doNotAutoCompressPictures"))
    if setting is None:
        setting = OxmlElement("w:doNotAutoCompressPictures")
        doc.settings.element.append(setting)
    setting.set(qn("w:val"), "true")
    target = save_revision(doc, HERE / "Lab08_revised.docx")
    check = Document(target)
    assert len(check.inline_shapes) == 4
    assert [canonical_xml(m) for m in check.element.body.xpath(".//m:oMath")] == math_before
    assert [canonical_xml(m) for m in check.element.body.xpath(".//w:object")] == objects_before
    text = "\n".join(p.text for p in check.paragraphs)
    assert "[Място за фигура" not in text
    assert not re.search(r"\b(?:9|10)\.\d+", text), "Old chapter numbering remains"
    assert sha(original) == checksum
    assert all(sha(DATA / r["path"]) == r["sha256"] for r in info["sources"])
    print("Saved", target.name, "legacy objects preserved:", len(objects_before))
    print(pd.DataFrame(metrics).to_string(index=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--reuse-results", action="store_true",
                        help="Keep the archived experiment exactly; renumber only")
    args = parser.parse_args()
    OUT.mkdir(exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10})
    diagram("fig8_1_pipeline", ["WAV записи → работно състояние, задание, идентификатор на запис",
            "Първи канал → пропускане на 1 s → прозорци 0,5 s без припокриване",
            "Изваждане на средното → 4 времеви + 7 спектрални признака",
            "Пет групови дяла: целият запис остава в една група",
            "Във всеки дял: StandardScaler само върху обучението → kNN / SVM; RF без мащабиране",
            "Предсказване върху задържаните записи → гласуване на прозорците",
            "Accuracy, Macro-F1 и матрици на ниво запис"])
    diagram("fig8_2_dataset", ["Mendeley Data: metadata.csv + data/ + scope_setup/ + quality/",
            "8 състояния: 4 работни състояния × без/с реверсиране",
            "Работна точка = състояние + задание за скорост",
            "current/ BIN • sound_vibrometer/ WAV • sound_phone/ M4A • vibration/ XLS",
            "Учебен анализ: WAV, 4 състояния без реверсиране, само общите задания",
            "Един запис на сензор и работна точка; сегментите не са повторни опити"])
    metrics, info = reuse_results() if args.reuse_results else experiment()
    revise(metrics, info)