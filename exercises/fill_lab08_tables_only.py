"""Fill result cells only, preserving every other DOCX byte and all figures."""
from pathlib import Path
import csv
import hashlib
import json
import re
import zipfile

from docx import Document

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "Lab08_figures_only.docx"


def main():
    source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    with (HERE / "figures_lab08/classification_metrics.csv").open(encoding="utf-8-sig", newline="") as stream:
        metrics = {(row["features"], row["model"]): row for row in csv.DictReader(stream)}
    doc = Document(SOURCE)
    assert len(doc.tables) == 3
    changes = {1: {}, 2: {}}
    for row_index, row in enumerate(doc.tables[1].rows[1:], 1):
        metric = metrics[("Времева и честотна област", row.cells[0].text.strip())]
        changes[1][(row_index, 1)] = f"{100 * float(metric['accuracy_record']):.1f}".replace(".", ",")
        changes[1][(row_index, 2)] = f"{float(metric['macro_f1_record']):.3f}".replace(".", ",")
    for row_index, row in enumerate(doc.tables[2].rows[1:], 1):
        for column, model in enumerate(("kNN", "SVM", "Random Forest"), 1):
            metric = metrics[(row.cells[0].text.strip(), model)]
            changes[2][(row_index, column)] = f"{float(metric['macro_f1_record']):.3f}".replace(".", ",")
    with zipfile.ZipFile(SOURCE) as archive:
        infos = archive.infolist()
        original = {item.filename: archive.read(item) for item in infos}
    xml = original["word/document.xml"].decode("utf-8")
    tables = list(re.finditer(r"<w:tbl\b[^>]*>.*?</w:tbl>", xml, re.DOTALL))
    assert len(tables) == 3
    replacements = []
    updated = xml
    for table_index, values in changes.items():
        table_xml = tables[table_index].group()
        rows = list(re.finditer(r"<w:tr\b[^>]*>.*?</w:tr>", table_xml, re.DOTALL))
        revised_table = table_xml
        for row_index in sorted({r for r, _ in values}, reverse=True):
            row_xml = rows[row_index].group()
            cells = list(re.finditer(r"<w:tc\b[^>]*>.*?</w:tc>", row_xml, re.DOTALL))
            revised_row = row_xml
            for column in sorted((c for r, c in values if r == row_index), reverse=True):
                cell = cells[column]
                assert doc.tables[table_index].cell(row_index, column).text.strip() == "—"
                revised_cell, count = re.subn(r"(<w:t\b[^>]*>)—(</w:t>)",
                                             lambda m: m[1] + values[(row_index, column)] + m[2], cell.group())
                assert count == 1
                revised_row = revised_row[:cell.start()] + revised_cell + revised_row[cell.end():]
            match = rows[row_index]
            revised_table = revised_table[:match.start()] + revised_row + revised_table[match.end():]
        assert updated.count(table_xml) == 1
        updated = updated.replace(table_xml, revised_table, 1)
        replacements.append((table_xml, revised_table))
    target = HERE / "Lab08_figures_and_tables.docx"
    version = 1
    while True:
        try:
            with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as archive:
                for item in infos:
                    archive.writestr(item, updated.encode("utf-8") if item.filename == "word/document.xml" else original[item.filename])
            break
        except PermissionError:
            version += 1
            target = HERE / f"Lab08_figures_and_tables_v{version}.docx"
    with zipfile.ZipFile(target) as archive:
        assert set(archive.namelist()) == set(original)
        restored = archive.read("word/document.xml").decode("utf-8")
        for before, after in replacements:
            assert restored.count(after) == 1
            restored = restored.replace(after, before, 1)
        assert restored == xml
        assert all(archive.read(name) == data for name, data in original.items() if name != "word/document.xml")
    check = Document(target)
    for table_index, values in changes.items():
        for (row, column), value in values.items():
            assert check.tables[table_index].cell(row, column).text.strip() == value
    assert len(check.inline_shapes) == len(doc.inline_shapes)
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == source_hash
    report = {"source": SOURCE.name, "output": target.name, "source_sha256": source_hash,
              "filled_result_cells": sum(map(len, changes.values())),
              "metrics": "recording-level out-of-fold; combined features for classification table; macro-F1 for feature comparison",
              "all_content_outside_two_result_tables_byte_identical": True,
              "all_other_package_parts_including_figures_byte_identical": True}
    (HERE / "tables_only_validation.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    print("Saved", target.name, ": 15 result cells filled; figures and all other content unchanged.")


if __name__ == "__main__":
    main()