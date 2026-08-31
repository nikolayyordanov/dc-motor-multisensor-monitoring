#!/usr/bin/env python3
"""Render the Markdown draft as an IEEE-style two-column Word document.

    python paper/md_to_docx_ieee.py                       # default draft -> .docx
    python paper/md_to_docx_ieee.py in.md out.docx

Layout follows the IEEE conference template: US Letter, 0.75"/1.0" top/bottom and
0.625" side margins, Times New Roman, a single-column title block followed by a
two-column body, roman-numbered small-caps section headings, "Fig. N." captions
below figures and "TABLE N" captions above tables.  Wide figures (aspect ratio
> 1.9) are placed in a full-width single-column band, mirroring ``figure*``.

Dependencies: python-docx, Pillow.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_IN = ROOT / "paper" / "multimodal_dual_task_draft.md"
DEFAULT_OUT = ROOT / "paper" / "Multimodal_DC_Motor_IEEE_draft.docx"

BODY_FONT = "Times New Roman"
COL_WIDTH_IN = 3.4
FULL_WIDTH_IN = 7.16
WIDE_ASPECT = 1.9

# --------------------------------------------------------------------------- #
# LaTeX -> unicode (light-weight: the draft only uses a small subset)
# --------------------------------------------------------------------------- #
_GREEK = {
    "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε", "lambda": "λ",
    "mu": "μ", "sigma": "σ", "tau": "τ", "phi": "φ", "omega": "ω", "Sigma": "Σ", "Delta": "Δ",
}
_SYMBOLS = {
    r"\times": "×", r"\cdot": "·", r"\pm": "±", r"\infty": "∞", r"\odot": "⊙",
    r"\top": "ᵀ", r"\sum": "Σ", r"\in": "∈", r"\approx": "≈", r"\leq": "≤", r"\geq": "≥",
    r"\quad": "  ", r"\qquad": "   ", r"\,": " ", r"\;": " ", r"\!": "", r"\\": "  ",
    r"\exp": "exp", r"\tanh": "tanh", r"\log": "log", r"\left": "", r"\right": "",
}
_SUP = str.maketrans("0123456789+-=()n2T", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿ²ᵀ")
_SUB = str.maketrans("0123456789+-=()aeijklmnoprstuvx",
                     "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ₐₑᵢⱼₖₗₘₙₒₚᵣₛₜᵤᵥₓ")
_ACCENTS = {"hat": "\u0302", "tilde": "\u0303", "bar": "\u0304", "vec": "\u20d7"}


def latex_to_text(src: str) -> str:
    """Best-effort conversion of the draft's LaTeX snippets to readable unicode."""
    s = src.strip()

    # font/emphasis wrappers carry no meaning here - unwrap them (may be nested)
    for _ in range(3):
        s = re.sub(r"\\(mathbf|mathrm|mathcal|mathit|text|textbf|boldsymbol)\{([^{}]*)\}", r"\2", s)
    # accents -> combining characters
    for _ in range(3):
        s = re.sub(r"\\(hat|tilde|bar|vec)\s*\{([^{}]*)\}",
                   lambda m: m.group(2) + _ACCENTS[m.group(1)], s)
        s = re.sub(r"\\(hat|tilde|bar|vec)\s+(\w)",
                   lambda m: m.group(2) + _ACCENTS[m.group(1)], s)

    s = re.sub(r"\\frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)", s)
    for name, ch in _GREEK.items():
        s = s.replace("\\" + name, ch)
    for pat, ch in _SYMBOLS.items():
        s = s.replace(pat, ch)

    def _script(m: re.Match, table) -> str:
        body = m.group(1) if m.group(1) is not None else (m.group(2) or "")
        if body and all(ord(c) in table for c in body):
            return body.translate(table)
        # not representable as unicode scripts: keep a flat, brace-free form
        return ("_" if table is _SUB else "^") + body

    s = re.sub(r"\^\{([^{}]*)\}|\^(\w)", lambda m: _script(m, _SUP), s)
    s = re.sub(r"_\{([^{}]*)\}|_(\w)", lambda m: _script(m, _SUB), s)
    s = s.replace("\\{", "{").replace("\\}", "}").replace("\\", "")
    return re.sub(r"\s+", " ", s).strip()


# --------------------------------------------------------------------------- #
# Word helpers
# --------------------------------------------------------------------------- #
def set_columns(section, num: int, space_in: float = 0.2) -> None:
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(int(space_in * 1440)))
    cols.set(qn("w:equalWidth"), "1")


def page_setup(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.625)
    section.right_margin = Inches(0.625)


def new_section(doc: Document, columns: int):
    section = doc.add_section(WD_SECTION.CONTINUOUS)
    page_setup(section)
    set_columns(section, columns)
    return section


def roman(n: int) -> str:
    vals = ((1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"),
            (50, "L"), (40, "XL"), (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I"))
    out = ""
    for v, sym in vals:
        while n >= v:
            out += sym
            n -= v
    return out


INLINE_RE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*|`[^`]+?`|\$[^$]+?\$|\[[^\]]+\]\([^)]+\))")


def add_runs(par, text: str, size: float, italic: bool = False, bold: bool = False) -> None:
    """Add inline-formatted runs (bold / italic / code / inline math / links)."""
    text = re.sub(r"\bFigure (\d+)", r"Fig. \1", text)
    for token in INLINE_RE.split(text):
        if not token:
            continue
        run_bold, run_italic, mono, body = bold, italic, False, token
        if token.startswith("**") and token.endswith("**"):
            run_bold, body = True, token[2:-2]
        elif token.startswith("*") and token.endswith("*"):
            run_italic, body = True, token[1:-1]
        elif token.startswith("`") and token.endswith("`"):
            mono, body = True, token[1:-1]
        elif token.startswith("$") and token.endswith("$"):
            run_italic, body = True, latex_to_text(token[1:-1])
        elif token.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            body = m.group(1) if m else token

        run = par.add_run(body)
        run.font.name = "Consolas" if mono else BODY_FONT
        run.font.size = Pt(size - 0.5 if mono else size)
        run.bold, run.italic = run_bold, run_italic


def body_par(doc, text: str, size: float = 10.0, indent: bool = True, space_after: float = 0.0):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    par.paragraph_format.space_after = Pt(space_after)
    par.paragraph_format.first_line_indent = Inches(0.2 if indent else 0)
    add_runs(par, text, size)
    return par


# --------------------------------------------------------------------------- #
# Markdown block parsing
# --------------------------------------------------------------------------- #
def split_blocks(lines: list[str]) -> list[tuple[str, object]]:
    """Group raw markdown lines into typed blocks."""
    blocks: list[tuple[str, object]] = []
    i, n = 0, len(lines)
    while i < n:
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped or stripped == "---":
            i += 1
            continue
        if stripped.startswith("```"):
            i, code = i + 1, []
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            blocks.append(("code", code))
            i += 1
            continue
        if stripped.startswith("$$"):
            i, math = i + 1, []
            while i < n and not lines[i].strip().startswith("$$"):
                math.append(lines[i].strip())
                i += 1
            blocks.append(("math", " ".join(math)))
            i += 1
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            blocks.append((f"h{level}", stripped.lstrip("#").strip()))
            i += 1
            continue
        if stripped.startswith("!["):
            m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
            if m:
                blocks.append(("image", m.group(2)))
                i += 1
                continue
        if stripped.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
            blocks.append(("table", rows))
            continue
        if stripped.startswith(">"):
            quote = []
            while i < n and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("quote", " ".join(quote)))
            continue
        if re.match(r"^[-*] ", stripped) or re.match(r"^\d+\. ", stripped):
            ordered = bool(re.match(r"^\d+\. ", stripped))
            items: list[str] = []
            while i < n and lines[i].strip():
                cur = lines[i].strip()
                if re.match(r"^[-*] ", cur) or re.match(r"^\d+\. ", cur):
                    items.append(re.sub(r"^([-*]|\d+\.)\s+", "", cur))
                elif items:
                    items[-1] += " " + cur
                else:
                    break
                i += 1
            blocks.append(("olist" if ordered else "ulist", items))
            continue

        para = []
        while i < n and lines[i].strip() and not re.match(r"^(\||>|#|!\[|```|\$\$|[-*] |\d+\. )", lines[i].strip()):
            para.append(lines[i].strip())
            i += 1
        if para:
            blocks.append(("para", " ".join(para)))
        else:
            i += 1
    return blocks


# --------------------------------------------------------------------------- #
# Renderers
# --------------------------------------------------------------------------- #
def add_figure(doc, img_path: Path, caption: str, fig_no: int, state: dict) -> None:
    with Image.open(img_path) as im:
        aspect = im.width / im.height
    wide = aspect > WIDE_ASPECT

    if wide and state["columns"] == 2:
        new_section(doc, 1)
        state["columns"] = 1
    width = FULL_WIDTH_IN if wide else COL_WIDTH_IN

    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(6)
    par.paragraph_format.space_after = Pt(3)
    par.add_run().add_picture(str(img_path), width=Inches(width))

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(caption) < 90 else WD_ALIGN_PARAGRAPH.JUSTIFY
    cap.paragraph_format.space_after = Pt(10)
    run = cap.add_run(f"Fig. {fig_no}. ")
    run.font.name, run.font.size = BODY_FONT, Pt(8)
    add_runs(cap, caption, 8.0)

    if wide:
        new_section(doc, 2)
        state["columns"] = 2


def add_table(doc, rows: list[list[str]], caption: str, table_no: int) -> None:
    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    head.paragraph_format.space_before = Pt(8)
    head.paragraph_format.space_after = Pt(0)
    run = head.add_run(f"TABLE {roman(table_no)}")
    run.font.name, run.font.size, run.font.small_caps = BODY_FONT, Pt(8), True

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(3)
    run = sub.add_run(re.sub(r"[*`$]", "", caption).upper())
    run.font.name, run.font.size, run.font.small_caps = BODY_FONT, Pt(8), True

    table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
    table.style = "Table Grid"
    table.autofit = True
    for r, row in enumerate(rows):
        for c, cell_text in enumerate(row):
            cell = table.cell(r, c)
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER if c else WD_ALIGN_PARAGRAPH.LEFT
            par.paragraph_format.space_after = Pt(0)
            add_runs(par, cell_text, 7.5, bold=(r == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def build(md_path: Path, out_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    blocks = split_blocks(lines)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    page_setup(doc.sections[0])
    set_columns(doc.sections[0], 1)

    state = {"columns": 1}
    sec_no = sub_no = fig_no = table_no = eq_no = 0
    pending_table_caption = ""
    pending_figure: Path | None = None
    started_body = False
    i = 0

    while i < len(blocks):
        kind, payload = blocks[i]

        # ---- title block ------------------------------------------------- #
        if kind == "h1":
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(10)
            run = par.add_run(str(payload))
            run.font.name, run.font.size = BODY_FONT, Pt(20)
            i += 1
            # author / affiliation lines until the first '## '
            while i < len(blocks) and blocks[i][0] == "para":
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.space_after = Pt(2)
                add_runs(par, str(blocks[i][1]).replace("*", ""), 10.5, italic=True)
                i += 1
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
            new_section(doc, 2)
            state["columns"] = 2
            continue

        # ---- headings ------------------------------------------------------ #
        if kind == "h2":
            title = str(payload)
            if title.lower().startswith("abstract"):
                i += 1
                text = str(blocks[i][1]) if i < len(blocks) and blocks[i][0] == "para" else ""
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                par.paragraph_format.space_after = Pt(6)
                run = par.add_run("Abstract—")
                run.bold = run.italic = True
                run.font.name, run.font.size = BODY_FONT, Pt(9)
                add_runs(par, text, 9.0, italic=True, bold=True)
                i += 1
                continue

            started_body = True
            sec_no, sub_no = sec_no + 1, 0
            clean = re.sub(r"^\d+\.\s*", "", title)
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_before = Pt(12)
            par.paragraph_format.space_after = Pt(4)
            run = par.add_run(f"{roman(sec_no)}.  {clean}")
            run.font.name, run.font.size, run.font.small_caps = BODY_FONT, Pt(10), True
            i += 1
            continue

        if kind == "h3":
            sub_no += 1
            clean = re.sub(r"^\d+\.\d+\s*", "", str(payload))
            par = doc.add_paragraph()
            par.paragraph_format.space_before = Pt(8)
            par.paragraph_format.space_after = Pt(3)
            run = par.add_run(f"{chr(64 + sub_no)}. {clean}")
            run.font.name, run.font.size, run.italic = BODY_FONT, Pt(10), True
            i += 1
            continue

        # ---- paragraphs and special-cased leading paragraphs --------------- #
        if kind == "para":
            text = str(payload)

            if text.startswith("**Keywords:**"):
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                par.paragraph_format.space_after = Pt(8)
                run = par.add_run("Index Terms—")
                run.bold = run.italic = True
                run.font.name, run.font.size = BODY_FONT, Pt(9)
                add_runs(par, text.replace("**Keywords:**", "").strip(), 9.0, italic=True)
                i += 1
                continue

            m_tab = re.match(r"\*\*Tables?\s*\d+\.\*\*\s*(.+)", text)
            if m_tab:
                pending_table_caption = m_tab.group(1)
                i += 1
                continue

            m_fig = re.match(r"\*\*(?:Figure|Fig\.)\s*(\d+)\.\*\*\s*(.+)", text, re.S)
            if m_fig and pending_figure is not None:
                add_figure(doc, pending_figure, m_fig.group(2), int(m_fig.group(1)), state)
                pending_figure = None
                i += 1
                continue

            body_par(doc, text, 10.0, indent=started_body, space_after=6)
            i += 1
            continue

        if kind == "image":
            path = (md_path.parent / str(payload)).resolve()
            if path.exists():
                fig_no += 1
                pending_figure = path
                # if the next block is not a caption, emit the figure immediately
                if not (i + 1 < len(blocks) and blocks[i + 1][0] == "para"
                        and re.match(r"\*\*(Figure|Fig\.)", str(blocks[i + 1][1]))):
                    add_figure(doc, path, "", fig_no, state)
                    pending_figure = None
            i += 1
            continue

        if kind == "table":
            table_no += 1
            add_table(doc, payload, pending_table_caption or "", table_no)
            pending_table_caption = ""
            i += 1
            continue

        if kind == "math":
            eq_no += 1
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_before = Pt(4)
            par.paragraph_format.space_after = Pt(6)
            run = par.add_run(latex_to_text(str(payload)))
            run.font.name, run.font.size, run.italic = BODY_FONT, Pt(10), True
            run = par.add_run(f"     ({eq_no})")
            run.font.name, run.font.size = BODY_FONT, Pt(10)
            i += 1
            continue

        if kind in ("ulist", "olist"):
            for k, item in enumerate(payload, start=1):
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                par.paragraph_format.left_indent = Inches(0.2)
                par.paragraph_format.space_after = Pt(2)
                bullet = f"{k}) " if kind == "olist" else "• "
                run = par.add_run(bullet)
                run.font.name, run.font.size = BODY_FONT, Pt(10)
                add_runs(par, item, 10.0)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            i += 1
            continue

        if kind == "code":
            for code_line in payload:
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Inches(0.12)
                par.paragraph_format.space_after = Pt(0)
                run = par.add_run(code_line or " ")
                run.font.name, run.font.size = "Consolas", Pt(7.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
            i += 1
            continue

        if kind == "quote":
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            par.paragraph_format.left_indent = Inches(0.15)
            par.paragraph_format.space_before = Pt(4)
            par.paragraph_format.space_after = Pt(6)
            add_runs(par, str(payload), 9.0, italic=True)
            for run in par.runs:
                run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
            i += 1
            continue

        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"Saved IEEE-style manuscript: {out_path}")


def main() -> None:
    md_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_IN
    out_path = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUT
    build(md_path, out_path)


if __name__ == "__main__":
    main()
